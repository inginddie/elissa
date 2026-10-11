"""Read-only file tools, confined to the folders the user allowed in settings."""
import fnmatch
import os
from datetime import datetime
from pathlib import Path
from typing import Callable

from . import ToolRegistry

CHUNK = 15000
SKIP_DIRS = {"node_modules", ".git", "__pycache__", ".venv", "venv", "AppData", "$Recycle.Bin", ".cache"}
SENSITIVE_DIRS = {".ssh", ".gnupg", ".aws", ".azure", ".kube", "AppData"}
SENSITIVE_NAMES = ["*.pem", "*.key", "*.pfx", "*.p12", "id_rsa*", "id_ed25519*", ".env", ".env.*", "*.kdbx",
                   "credentials*", "*.ovpn"]
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm", ".log", ".ini", ".yaml", ".yml", ".py",
            ".js", ".ts", ".sql", ".tsv", ".rtf"}
ALIASES = {
    "descargas": "Downloads", "downloads": "Downloads", "escritorio": "Desktop", "desktop": "Desktop",
    "documentos": "Documents", "documents": "Documents", "imagenes": "Pictures", "imágenes": "Pictures",
    "música": "Music", "musica": "Music", "videos": "Videos", "vídeos": "Videos",
}


def _human_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unit}"
        n /= 1024
    return f"{n:.1f} TB"


class FileTools:
    def __init__(self, allowed_folders: Callable[[], list[str]], handed_over: Callable[[], set[str]] = set):
        self.allowed_folders = allowed_folders
        # Files the user dropped on Elissa: allowed by that gesture even outside the allowed folders.
        self.handed_over = handed_over

    # ----- path safety -----

    def resolve(self, path: str) -> Path:
        raw = os.path.expandvars(os.path.expanduser((path or "~").strip().strip('"')))
        first = Path(raw).parts[0].lower() if Path(raw).parts else ""
        if first in ALIASES:
            raw = str(Path.home() / ALIASES[first] / Path(*Path(raw).parts[1:]))
        p = Path(raw)
        if not p.is_absolute():
            p = Path.home() / p
        p = p.resolve()
        if str(p).lower() in self.handed_over():
            if any(fnmatch.fnmatch(p.name.lower(), pat) for pat in SENSITIVE_NAMES):
                raise PermissionError("Ese archivo es sensible y no se puede leer.")
            return p
        roots = [Path(f).expanduser().resolve() for f in self.allowed_folders()]
        root = next((r for r in roots if p == r or p.is_relative_to(r)), None)
        if root is None:
            raise PermissionError(f"{p} está fuera de las carpetas permitidas ({', '.join(map(str, roots))}).")
        # Only what lies below the allowed root is checked: the user may explicitly allow e.g. a folder in AppData.
        below = p.relative_to(root).parts
        if any(part in SENSITIVE_DIRS for part in below) or any(fnmatch.fnmatch(p.name.lower(), pat)
                                                                for pat in SENSITIVE_NAMES):
            raise PermissionError("Ese archivo o carpeta es sensible y no se puede leer.")
        return p

    def _walk(self, folder: Path, max_entries: int = 60000):
        seen = 0
        for root, dirs, files in os.walk(folder):
            dirs[:] = [d for d in dirs if d not in SKIP_DIRS and d not in SENSITIVE_DIRS and not d.startswith(".")]
            for name in files:
                seen += 1
                if seen > max_entries:
                    return
                yield Path(root) / name

    # ----- tools -----

    def list_folder(self, path: str = "~", limit: int = 40) -> str:
        p = self.resolve(path)
        if not p.is_dir():
            raise NotADirectoryError(f"{p} no es una carpeta.")
        items = sorted(p.iterdir(), key=lambda x: x.stat().st_mtime if x.exists() else 0, reverse=True)
        lines = []
        for x in items[:limit]:
            try:
                st = x.stat()
            except OSError:
                continue
            kind = "carpeta" if x.is_dir() else _human_size(st.st_size)
            lines.append(f"{x.name}\t{kind}\t{datetime.fromtimestamp(st.st_mtime):%Y-%m-%d %H:%M}")
        more = f"\n…y {len(items) - limit} más" if len(items) > limit else ""
        return f"{p} (más recientes primero):\n" + "\n".join(lines) + more

    def recent_files(self, folder: str = "Descargas", extension: str = "", limit: int = 10) -> str:
        p = self.resolve(folder)
        ext = extension.lower().lstrip(".")
        files = [f for f in p.iterdir() if f.is_file() and (not ext or f.suffix.lower() == f".{ext}")]
        files.sort(key=lambda f: f.stat().st_mtime, reverse=True)
        if not files:
            return "No hay archivos que coincidan."
        return "\n".join(f"{f}\t{_human_size(f.stat().st_size)}\t{datetime.fromtimestamp(f.stat().st_mtime):%Y-%m-%d %H:%M}"
                         for f in files[:limit])

    def find_files(self, pattern: str, folder: str = "~", limit: int = 30) -> str:
        p = self.resolve(folder)
        pat = pattern.lower() if any(c in pattern for c in "*?") else f"*{pattern.lower()}*"
        hits = [f for f in self._walk(p) if fnmatch.fnmatch(f.name.lower(), pat)]
        hits.sort(key=lambda f: f.stat().st_mtime, reverse=True)
        if not hits:
            return f"No encontré archivos que coincidan con {pattern!r} en {p}."
        return "\n".join(f"{f}\t{datetime.fromtimestamp(f.stat().st_mtime):%Y-%m-%d}" for f in hits[:limit])

    def search_text(self, text: str, folder: str = "~", limit: int = 20) -> str:
        p = self.resolve(folder)
        needle = text.lower()
        out, scanned = [], 0
        for f in self._walk(p, 20000):
            if f.suffix.lower() not in TEXT_EXT:
                continue
            try:
                if f.stat().st_size > 2_000_000:
                    continue
                scanned += 1
                if scanned > 4000:
                    break
                for n, line in enumerate(f.read_text("utf-8", errors="ignore").splitlines(), 1):
                    if needle in line.lower():
                        out.append(f"{f}:{n}: {line.strip()[:200]}")
                        break
            except OSError:
                continue
            if len(out) >= limit:
                break
        return "\n".join(out) or f"No encontré {text!r} en archivos de texto de {p}."

    def read_file(self, path: str, start: int = 0) -> str:
        p = self.resolve(path)
        if not p.is_file():
            raise FileNotFoundError(f"No existe el archivo {p}.")
        text = extract_text(p)
        chunk = text[start:start + CHUNK]
        tail = (f"\n…(continúa: vuelve a llamar con start={start + CHUNK}; total {len(text)} caracteres)"
                if start + CHUNK < len(text) else "")
        return f"[{p.name}]\n{chunk}{tail}"


def extract_text(p: Path) -> str:
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(p))
        pages = [f"--- página {i} ---\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages, 1)]
        text = "\n".join(pages)
        return text if text.strip(" -\n0123456789páginas") else "(El PDF no tiene texto seleccionable; parece escaneado.)"
    if suffix == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for table in d.tables:
            parts += ["\t".join(c.text for c in row.cells) for row in table.rows]
        return "\n".join(parts)
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"--- hoja {ws.title} ---")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 300:
                    parts.append("…(más filas)")
                    break
                parts.append("\t".join("" if v is None else str(v) for v in row))
        return "\n".join(parts)
    if suffix in TEXT_EXT or suffix == "":
        raw = p.read_bytes()
        for enc in ("utf-8", "latin-1"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
    raise ValueError(f"No sé leer archivos {suffix or 'sin extensión'}.")


def register(reg: ToolRegistry, files: FileTools) -> None:
    reg.add("listar_carpeta", "Lista el contenido de una carpeta (más recientes primero). Acepta rutas o alias "
            "como 'Descargas', 'Escritorio', 'Documentos'.", "Revisando la carpeta…", files.list_folder,
            {"path": {"type": "string", "description": "Ruta o alias de la carpeta"},
             "limit": {"type": "integer", "description": "Máximo de elementos (por defecto 40)"}}, ["path"])
    reg.add("archivos_recientes", "Archivos más recientes de una carpeta, opcionalmente filtrados por extensión "
            "(ideal para 'el PDF que acabo de descargar').", "Buscando archivos recientes…", files.recent_files,
            {"folder": {"type": "string", "description": "Carpeta o alias, por defecto Descargas"},
             "extension": {"type": "string", "description": "p. ej. pdf, xlsx (opcional)"},
             "limit": {"type": "integer"}})
    reg.add("buscar_archivos", "Busca archivos por nombre (subcadena o comodines como *.pdf) dentro de una carpeta "
            "y sus subcarpetas.", "Buscando archivos…", files.find_files,
            {"pattern": {"type": "string"}, "folder": {"type": "string"}, "limit": {"type": "integer"}}, ["pattern"])
    reg.add("buscar_texto", "Busca un texto dentro de archivos de texto (txt, md, csv, json, código...) de una "
            "carpeta.", "Buscando dentro de archivos…", files.search_text,
            {"text": {"type": "string"}, "folder": {"type": "string"}, "limit": {"type": "integer"}}, ["text"])
    reg.add("leer_archivo", "Lee el contenido de un archivo: PDF, Word (.docx), Excel (.xlsx), CSV, texto o código. "
            "Para archivos largos usa 'start' para continuar.", "Leyendo el archivo…", files.read_file,
            {"path": {"type": "string"}, "start": {"type": "integer", "description": "Carácter desde el que leer"}},
            ["path"])
